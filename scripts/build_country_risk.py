"""Rebuild the packaged country risk table (fundamental_engine/brain/country_risk.json) from a saved Damodaran ctryprem folder.

    python -I scripts/build_country_risk.py runs/macro/ctryprem_<date> fundamental_engine/brain/country_risk.json

The folder holds ctryprem.html and ctryprem.xlsx saved verbatim by scripts/collect_market_context.py --country-risk (and its
MANIFEST.json, whose retrieval time is recorded). Country values come from the page (Country Risk Premium column), regional
GDP-weighted averages from the xlsx (Regional Weighted Averages sheet). Needs openpyxl (a maintenance dependency only; the
engine itself has none). ISO codes are typed here once (the page has none) so XBRL country members (country:TW) resolve.
Review the diff, run the tests and log the update in docs/CHANGELOG.md: a new table moves every country-risk row.
"""
import html
import json
import os
import re
import sys
import warnings
from datetime import datetime
import openpyxl
warnings.filterwarnings('ignore')
raw, out = sys.argv[1], sys.argv[2]
ISO2 = {'Albania': 'AL', 'Algeria': 'DZ', 'Andorra (Principality of)': 'AD', 'Angola': 'AO', 'Argentina': 'AR', 'Armenia': 'AM',
        'Aruba': 'AW', 'Australia': 'AU', 'Austria': 'AT', 'Azerbaijan': 'AZ', 'Bahamas': 'BS', 'Bahrain': 'BH', 'Bangladesh': 'BD',
        'Barbados': 'BB', 'Belarus': 'BY', 'Belgium': 'BE', 'Belize': 'BZ', 'Benin': 'BJ', 'Bermuda': 'BM', 'Bolivia': 'BO',
        'Bosnia and Herzegovina': 'BA', 'Botswana': 'BW', 'Brazil': 'BR', 'Brunei': 'BN', 'Bulgaria': 'BG', 'Burkina Faso': 'BF',
        'Cambodia': 'KH', 'Cameroon': 'CM', 'Canada': 'CA', 'Cape Verde': 'CV', 'Cayman Islands': 'KY', 'Chile': 'CL', 'China': 'CN',
        'Colombia': 'CO', 'Congo (Democratic Republic of)': 'CD', 'Congo (Republic of)': 'CG', 'Cook Islands': 'CK', 'Costa Rica': 'CR',
        "Côte d'Ivoire": 'CI', 'Croatia': 'HR', 'Cuba': 'CU', 'Curacao': 'CW', 'Cyprus': 'CY', 'Czech Republic': 'CZ', 'Denmark': 'DK',
        'Dominican Republic': 'DO', 'Ecuador': 'EC', 'Egypt': 'EG', 'El Salvador': 'SV', 'Estonia': 'EE', 'Ethiopia': 'ET', 'Fiji': 'FJ',
        'Finland': 'FI', 'France': 'FR', 'Gabon': 'GA', 'Gambia': 'GM', 'Georgia': 'GE', 'Germany': 'DE', 'Ghana': 'GH', 'Greece': 'GR',
        'Guatemala': 'GT', 'Guernsey (States of)': 'GG', 'Guinea': 'GN', 'Guinea-Bissau': 'GW', 'Guyana': 'GY', 'Haiti': 'HT',
        'Honduras': 'HN', 'Hong Kong': 'HK', 'Hungary': 'HU', 'Iceland': 'IS', 'India': 'IN', 'Indonesia': 'ID', 'Iran': 'IR', 'Iraq': 'IQ',
        'Ireland': 'IE', 'Isle of Man': 'IM', 'Israel': 'IL', 'Italy': 'IT', 'Jamaica': 'JM', 'Japan': 'JP', 'Jersey (States of)': 'JE',
        'Jordan': 'JO', 'Kazakhstan': 'KZ', 'Kenya': 'KE', 'Korea': 'KR', 'Korea, D.P.R.': 'KP', 'Kuwait': 'KW', 'Kyrgyzstan': 'KG',
        'Laos': 'LA', 'Latvia': 'LV', 'Lebanon': 'LB', 'Liberia': 'LR', 'Libya': 'LY', 'Liechtenstein': 'LI', 'Lithuania': 'LT',
        'Luxembourg': 'LU', 'Macao': 'MO', 'Macedonia': 'MK', 'Madagascar': 'MG', 'Malawi': 'MW', 'Malaysia': 'MY', 'Maldives': 'MV',
        'Mali': 'ML', 'Malta': 'MT', 'Mauritius': 'MU', 'Mexico': 'MX', 'Moldova': 'MD', 'Mongolia': 'MN', 'Montenegro': 'ME',
        'Montserrat': 'MS', 'Morocco': 'MA', 'Mozambique': 'MZ', 'Myanmar': 'MM', 'Namibia': 'NA', 'Nepal': 'NP', 'Netherlands': 'NL',
        'New Zealand': 'NZ', 'Nicaragua': 'NI', 'Niger': 'NE', 'Nigeria': 'NG', 'Norway': 'NO', 'Oman': 'OM', 'Pakistan': 'PK',
        'Panama': 'PA', 'Papua New Guinea': 'PG', 'Paraguay': 'PY', 'Peru': 'PE', 'Philippines': 'PH', 'Poland': 'PL', 'Portugal': 'PT',
        'Qatar': 'QA', 'Romania': 'RO', 'Russia': 'RU', 'Rwanda': 'RW', 'Saudi Arabia': 'SA', 'Senegal': 'SN', 'Serbia': 'RS',
        'Sierra Leone': 'SL', 'Singapore': 'SG', 'Slovakia': 'SK', 'Slovenia': 'SI', 'Solomon Islands': 'SB', 'Somalia': 'SO',
        'South Africa': 'ZA', 'Spain': 'ES', 'Sri Lanka': 'LK', 'St. Maarten': 'SX', 'St. Vincent & the Grenadines': 'VC', 'Sudan': 'SD',
        'Suriname': 'SR', 'Swaziland': 'SZ', 'Sweden': 'SE', 'Switzerland': 'CH', 'Syria': 'SY', 'Taiwan': 'TW', 'Tajikistan': 'TJ',
        'Tanzania': 'TZ', 'Thailand': 'TH', 'Togo': 'TG', 'Trinidad and Tobago': 'TT', 'Tunisia': 'TN', 'Turkey': 'TR',
        'Turks and Caicos Islands': 'TC', 'Uganda': 'UG', 'Ukraine': 'UA', 'United Arab Emirates': 'AE', 'United Kingdom': 'GB',
        'United States': 'US', 'Uruguay': 'UY', 'Uzbekistan': 'UZ', 'Venezuela': 'VE', 'Vietnam': 'VN', 'Yemen, Republic': 'YE',
        'Zambia': 'ZM', 'Zimbabwe': 'ZW'}
t = open(f'{raw}/ctryprem.html', encoding='latin-1').read()
updated = re.search(r'Last updated:\s*([A-Z][a-z]+ \d+, \d{4})', re.sub(r'<[^>]+>', ' ', t)).group(1)
t = re.sub(r'<style.*?</style>', '', t, flags=re.S)
pct = lambda s: round(float(s.rstrip('%')) / 100, 6) if s.endswith('%') else None
countries, notes = {}, {}
for r in re.findall(r'<tr[^>]*>(.*?)</tr>', t, flags=re.S | re.I):
    c = [html.unescape(re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', x))).strip() for x in re.findall(r'<td[^>]*>(.*?)</td>', r, flags=re.S | re.I)]
    if len(c) >= 5 and c[0] and c[3].endswith('%'):
        name = c[0]
        m = re.match(r'(.*?)\s*\((updated [^)]*)\)$', name)
        if m:
            name, notes[m.group(1)] = m.group(1), m.group(2)
        countries[name] = {'rating': c[1], 'default_spread': pct(c[2]), 'crp': pct(c[3]), 'erp': pct(c[4])}
wb = openpyxl.load_workbook(f'{raw}/ctryprem.xlsx', data_only=True, read_only=True)
rows = list(wb['ERPs by country'].iter_rows(values_only=True))
xlsx_note = ' '.join(str(x) for x in rows[1][:3] if x is not None)          # 'Date of update: 2026-01-01 (Updated Feb 16 for ...'
man = os.path.join(raw, 'MANIFEST.json')
retrieved = json.load(open(man))['retrieved_at_utc'][:10] if os.path.exists(man) else datetime.now().date().isoformat()
mature = rows[2][4]
us_erp = rows[3][4]
relvol = rows[5][4]
region_of = {r[0].split(' (updated')[0].strip(): r[1] for r in rows[8:] if r[0] and isinstance(r[1], str) and isinstance(r[5], (int, float))}
region_of.update({str(r[0]).strip(): r[1] for r in list(wb['Regional lookup table'].iter_rows(values_only=True))[1:] if r[0] and isinstance(r[1], str)})
for name, row in countries.items():
    if name in ISO2:
        row['iso2'] = ISO2[name]
    reg = region_of.get(name) or region_of.get(name + ' ')
    if reg:
        row['region'] = reg
    if name in notes:
        row['note'] = notes[name]
missing_iso = [n for n in countries if n not in ISO2]
rw = list(wb['Regional Weighted Averages'].iter_rows(values_only=True))
i = next(i for i, r in enumerate(rw) if r[0] == 'Region')
regions = {}
for r in rw[i + 1:]:
    if not r[0]:
        break
    regions[r[0]] = {'crp': round(r[2], 6), 'erp': round(r[1], 6), 'basis': 'GDP-weighted (Regional Weighted Averages)'}
j = next(i for i, r in enumerate(rw) if r[0] == 'For updating industry average spreadsheets')
for r in rw[j + 2:]:
    if not r[0]:
        break
    if r[0] not in regions and r[0] not in ('US', 'China', 'Japan', 'India'):     # single countries: the country rows
        regions[r[0]] = {'crp': round(r[1] - mature, 6), 'erp': round(r[1], 6), 'basis': 'GDP-weighted grouping (ERP - mature market premium)'}
k = next(i for i, r in enumerate(rw) if r[0] == 'For other regional groupings')
for r in rw[k + 2:]:
    if not r[0]:
        break
    regions[r[0]] = {'crp': round(r[1] - mature, 6), 'erp': round(r[1], 6), 'basis': 'GDP-weighted grouping (ERP - mature market premium)'}
data = {
    'title': 'Damodaran, Country Default Spreads and Risk Premiums (equity country risk premium = rating-based default spread x '
             'relative equity market volatility)',
    'url': 'https://pages.stern.nyu.edu/~adamodar/New_Home_Page/datafile/ctryprem.html',
    'xlsx_url': 'https://pages.stern.nyu.edu/~adamodar/pc/datasets/ctryprem.xlsx',
    'as_of': datetime.strptime(updated, '%B %d, %Y').date().isoformat(), 'page_last_updated': updated, 'retrieved_at': retrieved,
    'xlsx_date_note': xlsx_note,
    'mature_market_erp': mature, 'us_erp': us_erp, 'relative_equity_volatility': round(relvol, 6),
    'base_country': 'United States',
    'base_note': ('the engine discounts with the US implied ERP (Damodaran\'s S&P 500 implied premium), which already carries the US '
                  'country risk premium after the May 2025 Moody\'s downgrade to Aa1 (US ERP = mature premium + US CRP); so only the part '
                  'of a country\'s CRP above the US CRP is added'),
    'countries': countries, 'regions': regions,
    'country_values_from': 'the HTML page (Country Risk Premium column, two decimals of a percent)',
    'region_values_from': 'the xlsx (Regional Weighted Averages sheet, GDP weights 2024); grouping CRP = grouping ERP - mature premium',
}
txt = json.dumps({k: v for k, v in data.items() if k not in ('countries', 'regions')}, indent=1, ensure_ascii=False)[:-2]
txt += ',\n "countries": {\n' + ',\n'.join(f'  {json.dumps(k, ensure_ascii=False)}: {json.dumps(v, ensure_ascii=False)}' for k, v in countries.items())
txt += '\n },\n "regions": {\n' + ',\n'.join(f'  {json.dumps(k, ensure_ascii=False)}: {json.dumps(v, ensure_ascii=False)}' for k, v in regions.items()) + '\n }\n}\n'
assert json.loads(txt) == data
open(out, 'w', encoding='utf-8').write(txt)
print(len(countries), 'countries;', len(regions), 'regions; no iso2:', missing_iso, 'mature', mature, 'us', us_erp, 'relvol', relvol, updated)
print({k: v['crp'] for k, v in regions.items()})
print('no region:', [n for n, r in countries.items() if 'region' not in r]); print(countries['United States'], countries['Taiwan'], countries['Turkey'])
